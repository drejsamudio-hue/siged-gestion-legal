"""Extractor de texto: PDF (PyMuPDF), .txt/.md (lectura directa),
.docx (python-docx), .doc (antiword o lectura binaria).

`extraer()` maneja solo PDFs (compatibilidad con el pipeline existente).
`extraer_universal()` despacha por extension y cubre todos los formatos.
"""
from __future__ import annotations

import hashlib
import logging
import os
import subprocess
from pathlib import Path

import fitz  # PyMuPDF

logging.disable(logging.WARNING)

PAGS_INICIO = 6
PAGS_FINAL = 1

_TEXTO_PLANO = {".txt", ".md", ".csv", ".log", ".json"}
_DOCX = {".docx"}
_DOC = {".doc"}
_PDF = {".pdf"}

# OCR para PDF escaneados (Tesseract con espanol). Solo las primeras paginas: alcanza para clasificar.
from config import OCR_CACHE, TESSDATA  # noqa: E402,F401
OCR_PAGINAS = 3
OCR_DPI = 300
OCR_MINIMO = 100
OCR_CALIDAD_MINIMA = 0.06
_PALABRAS_FRECUENTES = set(
    "de la el en y a los las del que por con se su al para una un es lo como sr juez expte autos parte no".split())


def extraer(path: Path, max_chars: int = 40000) -> tuple[str | None, str | None]:
    try:
        doc = fitz.open(str(path))
    except Exception as exc:  # noqa: BLE001
        return None, f"ERROR_APERTURA: {exc}"
    try:
        if doc.is_encrypted and not doc.authenticate(""):
            return None, "ERROR_CIFRADO"
        n = doc.page_count
        indices = list(range(min(PAGS_INICIO, n)))
        for i in range(max(0, n - PAGS_FINAL), n):
            if i not in indices:
                indices.append(i)
        partes = []
        total = 0
        for i in indices:
            try:
                t = doc.load_page(i).get_text()
            except Exception:  # noqa: BLE001
                continue
            partes.append(t)
            total += len(t)
            if total >= max_chars:
                break
        return "\n".join(partes), None
    except Exception as exc:  # noqa: BLE001
        return None, f"ERROR_LECTURA: {exc}"
    finally:
        try:
            doc.close()
        except Exception:  # noqa: BLE001
            pass


def paginas_pdf(path: Path) -> int:
    try:
        with fitz.open(str(path)) as doc:
            return doc.page_count
    except Exception:  # noqa: BLE001
        return 0


def _extraer_texto_plano(path: Path, max_chars: int) -> tuple[str | None, str | None]:
    try:
        raw = path.read_bytes()
    except OSError as exc:
        return None, f"ERROR_LECTURA: {exc}"
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(enc)[:max_chars], None
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")[:max_chars], None


def _extraer_docx(path: Path, max_chars: int) -> tuple[str | None, str | None]:
    try:
        from docx import Document
        doc = Document(str(path))
        partes = []
        for sec in doc.sections:
            partes += [p.text for p in sec.header.paragraphs]
        partes += [p.text for p in doc.paragraphs]
        for tabla in doc.tables:
            for fila in tabla.rows:
                partes.append(" | ".join(c.text for c in fila.cells))
        texto = "\n".join(partes)[:max_chars]
        return texto if texto.strip() else None, None
    except ImportError:
        return None, "python-docx_no_instalado"
    except Exception as exc:  # noqa: BLE001
        return None, f"ERROR_DOCX: {exc}"


def _extraer_doc(path: Path, max_chars: int) -> tuple[str | None, str | None]:
    try:
        result = subprocess.run(
            ["antiword", str(path)],
            capture_output=True, text=True, timeout=30,
        )
        if result.returncode == 0 and result.stdout.strip():
            return result.stdout[:max_chars], None
    except FileNotFoundError:
        pass
    except Exception:  # noqa: BLE001
        pass
    try:
        raw = path.read_bytes()[:max_chars * 2]
        texto = raw.decode("latin-1", errors="replace")
        import re
        texto = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", " ", texto)
        if len(texto.strip()) > 50:
            return texto[:max_chars], "doc_lectura_binaria"
    except Exception:  # noqa: BLE001
        pass
    return None, "ERROR_DOC_sin_antiword"


def ocr_disponible() -> bool:
    return (TESSDATA / "spa.traineddata").is_file()


def _clave_ocr(path: Path) -> str:
    st = path.stat()
    return hashlib.sha1(f"{path.name}|{st.st_size}|{int(st.st_mtime)}|{OCR_DPI}|{OCR_PAGINAS}".encode()).hexdigest()


def calidad_texto(texto: str) -> float:
    """Proporcion de palabras muy frecuentes del castellano: el OCR ilegible casi no las tiene."""
    import re
    palabras = re.findall(r"[a-záéíóúñ]+", texto.lower())
    return sum(len(w) >= 2 and w in _PALABRAS_FRECUENTES for w in palabras) / max(1, len(palabras))


def ocr_pdf(path: Path, paginas: int = OCR_PAGINAS) -> str | None:
    """Texto por OCR de las primeras paginas; se guarda en cache para no repetirlo en cada simulacion."""
    if not ocr_disponible():
        return None
    try:
        cache = OCR_CACHE / f"{_clave_ocr(path)}.txt"
        if cache.is_file():
            return cache.read_text(encoding="utf-8")
        with fitz.open(str(path)) as doc:
            partes = []
            for i in range(min(paginas, doc.page_count)):
                pg = doc.load_page(i)
                tp = pg.get_textpage_ocr(language="spa", tessdata=str(TESSDATA), full=True, dpi=OCR_DPI)
                partes.append(pg.get_text(textpage=tp))
        texto = "\n".join(partes)
        OCR_CACHE.mkdir(parents=True, exist_ok=True)
        cache.write_text(texto, encoding="utf-8")
        return texto
    except Exception:  # noqa: BLE001
        return None


def extraer_universal(path: Path, max_chars: int = 40000) -> tuple[str | None, str | None]:
    ext = path.suffix.lower()
    if ext in _PDF:
        texto, err = extraer(path, max_chars)
        if err is None and len((texto or "").strip()) < OCR_MINIMO:
            ocr = ocr_pdf(path)
            if ocr and len(ocr.strip()) >= OCR_MINIMO:
                if calidad_texto(ocr) < OCR_CALIDAD_MINIMA:
                    return None, "ocr_ilegible"
                return ocr[:max_chars], "ocr"
        return texto, err
    if ext in _TEXTO_PLANO:
        return _extraer_texto_plano(path, max_chars)
    if ext in _DOCX:
        return _extraer_docx(path, max_chars)
    if ext in _DOC:
        return _extraer_doc(path, max_chars)
    return None, f"formato_no_soportado:{ext}"
